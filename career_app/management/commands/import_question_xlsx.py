"""
Imports the multi-sheet interview question workbook into the FAQQuestion table.

The workbook (data/*.xlsx) has one sheet per industry, each with the columns:

    topic | subtopic | difficulty level | questions

FAQQuestion has no `subtopic` column, and interview selection matches questions
to a candidate's resume SKILLS via FAQQuestion.topic (see _matching_bank_topics
in resume_utils.py). The specific `subtopic` (e.g. "supervised learning",
"customer support") is the skill-like value that matches a resume, not the broad
industry in `topic` — so subtopic becomes FAQQuestion.topic. The difficulty
labels ("Fresher", "1-3 Years", "3-5 Years", "5+ Years") are stored verbatim;
DIFFICULTY_DB_KEYS already maps them onto Easy/Intermediate/Hard.

Dedup is handled by the (topic, question_text) unique constraint via
ignore_conflicts, so re-running only inserts new rows.

Usage:
    python manage.py import_question_xlsx
    python manage.py import_question_xlsx --path "data/Some Workbook.xlsx"
"""

import glob
import os

from django.conf import settings
from django.core.cache import cache
from django.core.management.base import BaseCommand

from career_app.models import FAQQuestion
from career_app.resume_utils import ACADEMIC_TOPIC, BEHAVIOURAL_TOPIC, EXPERIENCE_TOPIC

# HR-bank workbooks (e.g. data/behavioural_experience_questions.xlsx) have a
# `Category` column instead of topic/subtopic. Their rows go into the fixed bank
# topics the interview picks Q1-10 from, so they never mix with domain topics.
CATEGORY_TOPICS = {
    "behavioral": BEHAVIOURAL_TOPIC,
    "behavioural": BEHAVIOURAL_TOPIC,
    "experience": EXPERIENCE_TOPIC,
    "academic": ACADEMIC_TOPIC,
    "fresher": ACADEMIC_TOPIC,
}


def _find_col(header, *names):
    """Index of the first header cell matching any of `names` (case-insensitive)."""
    lowered = [str(h).strip().lower() if h is not None else "" for h in header]
    for name in names:
        if name in lowered:
            return lowered.index(name)
    return None


class Command(BaseCommand):
    help = "Load the multi-sheet .xlsx question workbook into FAQQuestion (topic = subtopic)."

    def add_arguments(self, parser):
        parser.add_argument(
            "--path",
            default=None,
            help="Path to one .xlsx (defaults to every data/*.xlsx).",
        )

    def handle(self, *args, **options):
        # Every workbook under data/ by default — the domain bank and the HR bank
        # are separate files, and picking only the newest one silently skipped
        # whichever was older. Re-importing is safe (duplicates are ignored).
        if options["path"]:
            paths = [options["path"]]
        else:
            paths = sorted(glob.glob(os.path.join(str(settings.BASE_DIR), "data", "*.xlsx")))
            if not paths:
                self.stderr.write(self.style.ERROR("No .xlsx found under data/."))
                return

        for path in paths:
            if not os.path.exists(path):
                self.stderr.write(self.style.ERROR(f"File not found: {path}"))
                continue
            self._import_workbook(path)

    def _import_workbook(self, path):
        import openpyxl  # local import so the app runs without openpyxl until this command is used

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

        to_create = []
        skipped = 0
        seen_keys = set()  # (topic, question_text) already queued this run
        per_sheet = {}

        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            try:
                header = next(rows)
            except StopIteration:
                continue

            i_topic = _find_col(header, "subtopic")          # skill-level value → FAQQuestion.topic
            i_industry = _find_col(header, "topic")           # broad industry → FAQQuestion.industry
            i_diff = _find_col(header, "difficulty level", "difficulty", "level")
            i_q = _find_col(header, "questions", "question_text", "question")
            i_category = _find_col(header, "category")        # HR-bank workbooks only
            if i_q is None or (i_topic is None and i_industry is None and i_category is None):
                skipped += ws.max_row or 0
                continue

            count = 0
            for row in rows:
                def cell(idx):
                    return row[idx] if idx is not None and idx < len(row) else None

                if i_topic is None and i_industry is None:
                    # HR bank: an unknown category is skipped rather than guessed.
                    industry = ""
                    topic = CATEGORY_TOPICS.get(str(cell(i_category) or "").strip().lower(), "")
                else:
                    industry = str(cell(i_industry) or "").strip().lower()
                    topic = str(cell(i_topic) or "").strip().lower() or industry
                question_text = str(cell(i_q) or "").strip()
                difficulty = str(cell(i_diff) or "").strip() or None

                if not topic or not question_text:
                    skipped += 1
                    continue

                key = (topic, question_text[:500])
                if key in seen_keys:
                    skipped += 1
                    continue
                seen_keys.add(key)

                to_create.append(
                    FAQQuestion(topic=topic, industry=industry,
                                question_text=question_text[:500], difficulty=difficulty)
                )
                count += 1
            if count:
                per_sheet[ws.title] = count

        if not to_create:
            self.stderr.write(self.style.WARNING("No valid rows found in workbook."))
            return

        before = FAQQuestion.objects.count()
        FAQQuestion.objects.bulk_create(to_create, ignore_conflicts=True, batch_size=1000)
        after = FAQQuestion.objects.count()
        cache.delete('faq_domain_topics')  # _matching_bank_topics caches the topic list

        # bulk_create(ignore_conflicts) skips rows that already existed from a
        # previous run, leaving their `industry` blank. Backfill it per subtopic
        # so re-running this command upgrades an older import in place.
        topic_industry = {}
        for fq in to_create:
            if fq.industry:
                topic_industry.setdefault(fq.topic, fq.industry)
        filled = 0
        for topic, ind in topic_industry.items():
            filled += FAQQuestion.objects.filter(topic=topic, industry='').update(industry=ind)
        if filled:
            self.stdout.write(f"Backfilled industry on {filled} existing row(s).")

        for title, n in per_sheet.items():
            self.stdout.write(f"  {title}: {n} row(s) read")
        self.stdout.write(self.style.SUCCESS(
            f"Imported {after - before} new question(s) from {os.path.basename(path)} "
            f"({skipped} row(s) skipped, {after} total in DB)."
        ))
